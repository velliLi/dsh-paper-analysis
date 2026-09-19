#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx_core —— markdown → Word 渲染核心（三种版本）。

本文件是技能自带的可移植核心：不依赖 DSH，只依赖 python-docx。
CLI 见同目录的 make_docx.py。

三种版本（kind）：
  detail  详细版：封面、自动目录域、页脚页码、表格样式、引文块、图表自动编号
  brief   简略版：2–3 页整段文字，无封面无标题无目录
  notes   摘录版：只有"论文在做什么 / 创新点 / 该论文提供的思路"三节，
          篇幅最小，不含任何评级、证据链标记与编辑判断

排印（三版一致，用户指定）：中文宋体小四（12 pt）、英文 Times New Roman、1.5 倍行距、全文黑色。
"""

from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, field

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ModuleNotFoundError:  # pragma: no cover
    raise SystemExit(
        "python-docx 未安装。请先执行：pip install python-docx\n"
        "（生成 Word 报告依赖它；装不了时请把 markdown 直接交给使用者）"
    )

# ── design tokens ───────────────────────────────────────────────────────────
FONT_CJK = "宋体"
FONT_LATIN = "Times New Roman"
FONT_HEI = "黑体"
BODY_SIZE = 12.0          # 小四
LINE_SPACING = 1.5        # 1.5 倍行距
BLACK = RGBColor(0x00, 0x00, 0x00)   # 全文黑色（用户要求）

KINDS = ("detail", "brief", "notes")


@dataclass
class DocStyle:
    kind: str
    base_font: str = FONT_CJK
    latin_font: str = FONT_LATIN
    base_size: float = BODY_SIZE
    line_spacing: float = LINE_SPACING
    table_size: float = 10.5
    code_size: float = 10.0
    h1_size: float = 14.0
    h2_size: float = 12.0
    h_space_before: float = 12.0
    body_space_after: float = 6.0
    page_margin_cm: float = 2.5
    add_toc: bool = True
    add_cover: bool = True
    image_width_cm: float = 15.0

    @staticmethod
    def for_kind(kind: str) -> "DocStyle":
        if kind == "brief":
            return DocStyle(kind="brief", page_margin_cm=2.2, body_space_after=4.0,
                            h_space_before=8.0, add_toc=False, add_cover=False,
                            image_width_cm=13.0)
        if kind == "notes":
            # 最小篇幅：无封面、无标题、无目录，页边距与段间距都收紧
            return DocStyle(kind="notes", page_margin_cm=2.0, body_space_after=3.0,
                            h_space_before=7.0, add_toc=False, add_cover=False,
                            image_width_cm=12.0)
        return DocStyle(kind="detail")


@dataclass
class Meta:
    title: str = ""
    subtitle: str = ""
    author: str = ""
    date: str = ""
    paper: str = ""
    source: str = ""
    version_label: str = ""
    extra: dict = field(default_factory=dict)


# ── markdown → block AST ────────────────────────────────────────────────────
META_RE = re.compile(r"^<!--\s*meta:\s*(\{.*\})\s*-->\s*$", re.DOTALL)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
ULIST_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
OLIST_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
QUOTE_RE = re.compile(r"^>\s?(.*)$")
FENCE_RE = re.compile(r"^```(.*)$")
IMG_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
HR_RE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
INLINE_RE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`|~~.+?~~)"
)


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [p.replace("\\|", "|").strip() for p in re.split(r"(?<!\\)\|", line)]


def parse_markdown(text: str) -> tuple[Meta, list[tuple[str, object]]]:
    meta = Meta()
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[tuple[str, object]] = []
    i, n = 0, len(lines)
    para_buf: list[str] = []

    def flush_para() -> None:
        if para_buf:
            joined = " ".join(s.strip() for s in para_buf).strip()
            if joined:
                blocks.append(("para", joined))
            para_buf.clear()

    while i < n:
        raw = lines[i]
        stripped = raw.strip()

        m = META_RE.match(stripped)
        if m and not blocks:
            try:
                data = json.loads(m.group(1))
                meta.title = str(data.get("title", ""))
                meta.subtitle = str(data.get("subtitle", ""))
                meta.author = str(data.get("author", ""))
                meta.date = str(data.get("date", ""))
                meta.paper = str(data.get("paper", ""))
                meta.source = str(data.get("source", ""))
                meta.version_label = str(data.get("version_label", ""))
                meta.extra = data
            except json.JSONDecodeError:
                meta.extra = {"meta_parse_error": True}
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        if FENCE_RE.match(stripped):
            flush_para()
            lang = FENCE_RE.match(stripped).group(1).strip()  # type: ignore[union-attr]
            i += 1
            body: list[str] = []
            while i < n and not FENCE_RE.match(lines[i].strip()):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", (lang, "\n".join(body))))
            continue

        if HR_RE.match(stripped):
            flush_para()
            blocks.append(("hr", None))
            i += 1
            continue

        hm = HEADING_RE.match(stripped)
        if hm:
            flush_para()
            blocks.append(("heading", (len(hm.group(1)), hm.group(2).strip())))
            i += 1
            continue

        im = IMG_RE.match(stripped)
        if im:
            flush_para()
            blocks.append(("image", (im.group(1), im.group(2))))
            i += 1
            continue

        if "|" in stripped and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip()):
            flush_para()
            header = split_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i].strip()))
                i += 1
            blocks.append(("table", (header, rows)))
            continue

        if QUOTE_RE.match(stripped):
            flush_para()
            body = []
            while i < n and QUOTE_RE.match(lines[i].strip()):
                body.append(QUOTE_RE.match(lines[i].strip()).group(1))  # type: ignore[union-attr]
                i += 1
            blocks.append(("quote", "\n".join(body).strip()))
            continue

        ul = ULIST_RE.match(raw)
        if ul:
            flush_para()
            items = []
            while i < n:
                um = ULIST_RE.match(lines[i])
                if not um:
                    break
                items.append((len(um.group(1)) // 2, um.group(2).strip()))
                i += 1
            blocks.append(("ulist", items))
            continue

        ol = OLIST_RE.match(raw)
        if ol:
            flush_para()
            items = []
            while i < n:
                om = OLIST_RE.match(lines[i])
                if not om:
                    break
                items.append((len(om.group(1)) // 2, om.group(3).strip()))
                i += 1
            blocks.append(("olist", items))
            continue

        para_buf.append(raw)
        i += 1

    flush_para()
    return meta, blocks


# ── fonts & runs ────────────────────────────────────────────────────────────
def _bind(run, cjk_font: str = FONT_CJK, latin_font: str = FONT_LATIN) -> None:
    """一个 run 上同时绑定中文与西文字体（Word 按字符脚本自动取用）。"""
    run.font.name = latin_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin_font)
    rfonts.set(qn("w:hAnsi"), latin_font)
    rfonts.set(qn("w:cs"), latin_font)
    rfonts.set(qn("w:eastAsia"), cjk_font)


def _bind_style(style, cjk_font: str = FONT_CJK, latin_font: str = FONT_LATIN) -> None:
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin_font)
    rfonts.set(qn("w:hAnsi"), latin_font)
    rfonts.set(qn("w:cs"), latin_font)
    rfonts.set(qn("w:eastAsia"), cjk_font)


def _run(paragraph, text: str, size: float, *, bold=False, italic=False,
         strike=False, color=None, cjk_font=None, latin_font=None) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    _bind(run, cjk_font or FONT_CJK, latin_font or FONT_LATIN)
    if strike:
        run.font.strike = True
    if color is not None:
        run.font.color.rgb = color


def add_runs(paragraph, text: str, size: float) -> None:
    """把行内 markdown（粗体 / 斜体 / 等宽 / 删除线）渲染成 run。"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos : m.start()], size)
        token = m.group(0)
        if token.startswith("***") and token.endswith("***"):
            _run(paragraph, token[3:-3], size, bold=True, italic=True)
        elif token.startswith("**") and token.endswith("**"):
            _run(paragraph, token[2:-2], size, bold=True, color=BLACK)
        elif token.startswith("`") and token.endswith("`"):
            _run(paragraph, token[1:-1], size - 0.5)
        elif token.startswith("~~") and token.endswith("~~"):
            _run(paragraph, token[2:-2], size, strike=True)
        elif token.startswith("*") and token.endswith("*"):
            _run(paragraph, token[1:-1], size, italic=True)
        else:
            _run(paragraph, token, size)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size)


def _shade_paragraph(paragraph, fill: str) -> None:
    pr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _shade_cell(cell, fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._element.append(el)


def _refresh_fields_on_open(document) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is not None:
        settings.remove(existing)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


# ── document construction ───────────────────────────────────────────────────
def _set_default_style(document, style: DocStyle) -> None:
    normal = document.styles["Normal"]
    normal.font.name = style.latin_font
    normal.font.size = Pt(style.base_size)
    _bind_style(normal, style.base_font, style.latin_font)
    pf = normal.paragraph_format
    pf.line_spacing = style.line_spacing
    pf.space_after = Pt(style.body_space_after)

    for name, size, before in (
        ("Heading 1", style.h1_size, style.h_space_before),
        ("Heading 2", style.h2_size, style.h_space_before * 0.7),
        ("Heading 3", style.base_size, style.h_space_before * 0.6),
        ("Heading 4", style.base_size, style.h_space_before * 0.5),
    ):
        st = document.styles[name]
        st.font.name = style.latin_font
        st.font.size = Pt(size)
        st.font.bold = True
        _bind_style(st, FONT_HEI, style.latin_font)
        # python-docx 模板里的标题样式自带主题蓝，用户要求全文黑色：把颜色也压成黑
        st.font.color.rgb = BLACK
        rpr = st.element.get_or_add_rPr()
        for color_el in rpr.findall(qn("w:color")):
            color_el.set(qn("w:val"), "000000")
            for theme_attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
                if color_el.get(qn(theme_attr)) is not None:
                    del color_el.attrib[qn(theme_attr)]
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(3 if style.kind in ("brief", "notes") else 5)
        st.paragraph_format.line_spacing = style.line_spacing
        st.paragraph_format.keep_with_next = True


def _page_setup(document, style: DocStyle, footer_left: str) -> None:
    sec = document.sections[0]
    sec.left_margin = Cm(style.page_margin_cm)
    sec.right_margin = Cm(style.page_margin_cm)
    sec.top_margin = Cm(style.page_margin_cm)
    sec.bottom_margin = Cm(style.page_margin_cm)
    footer = sec.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = (footer_left or "").strip()
    if label:
        _run(p, label + "　", 8)
    _run(p, "第 ", 8)
    _field(p, "PAGE", "1")
    _run(p, " / ", 8)
    _field(p, "NUMPAGES", "1")
    _run(p, " 页", 8)


def _cover(document, meta: Meta) -> None:
    for _ in range(3):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, meta.title or "论文分析报告", 20, bold=True, color=BLACK)
    if meta.subtitle:
        q = document.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(q, meta.subtitle, 12, color=BLACK)
    document.add_paragraph()
    if meta.paper:
        for line in meta.paper.split("\n"):
            r = document.add_paragraph()
            r.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(r, line.strip(), 11)
    document.add_paragraph()
    for label, value in (("作者", meta.author), ("日期", meta.date), ("来源", meta.source)):
        if value:
            r = document.add_paragraph()
            r.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(r, f"{label}：{value}", 10.5)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_toc(document) -> None:
    h = document.add_paragraph()
    _run(h, "目录", 14, bold=True, color=BLACK)
    p = document.add_paragraph()
    _field(p, 'TOC \\o "1-3" \\h \\z \\u', "（在 Word 中打开后目录会自动刷新；若未刷新请按 Ctrl+A 再按 F9）")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_table(document, header: list[str], rows: list[list[str]], style: DocStyle) -> None:
    cols = max([len(header)] + [len(r) for r in rows]) if (header or rows) else 0
    if cols == 0:
        return
    table = document.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for idx in range(cols):
        cell = hdr[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = style.line_spacing
        add_runs(p, (header[idx] if idx < len(header) else "") or " ", style.table_size)
        for run in p.runs:
            run.bold = True
        _shade_cell(cell, "DCE3F0")
    for r in rows:
        cells = table.add_row().cells
        for idx in range(cols):
            cell = cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = style.line_spacing
            add_runs(p, r[idx] if idx < len(r) else "", style.table_size)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_image(document, caption: str, path: str, style: DocStyle, base_dir: str = "") -> None:
    resolved = path.strip()
    if not os.path.isabs(resolved):
        candidates = [os.path.abspath(resolved)]
        if base_dir:
            candidates.insert(0, os.path.abspath(os.path.join(base_dir, resolved)))
        resolved = next((c for c in candidates if os.path.exists(c)), candidates[0])
    if not os.path.exists(resolved):
        p = document.add_paragraph()
        _run(p, f"[插图缺失：{path}（{caption}）]", style.base_size, italic=True)
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    try:
        run.add_picture(resolved, width=Cm(style.image_width_cm))
    except Exception as exc:
        _run(p, f"[插图无法嵌入：{path}（{exc}）]", style.base_size, italic=True)
        return
    if caption:
        c = document.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(c, caption, style.table_size)


def _save_with_retry(document, out_path: str, attempts: int = 4) -> None:
    """Word 可能仍占着上一次的文件：短暂重试而不是让整轮失败。"""
    import time

    last: Exception | None = None
    for attempt in range(attempts):
        try:
            document.save(out_path)
            return
        except PermissionError as exc:
            last = exc
            time.sleep(0.8 * (attempt + 1))
    raise SystemExit(f"无法写入 {out_path}：文件被占用（Word 或预览窗口正打开它）。关闭后重试。\n底层错误：{last}")


def plain_text(blocks: list[tuple[str, object]]) -> str:
    parts: list[str] = []
    for tag, payload in blocks:
        if tag in ("para", "quote"):
            parts.append(str(payload))
        elif tag == "heading":
            parts.append(str(payload[1]))  # type: ignore[index]
        elif tag in ("ulist", "olist"):
            parts.extend(item for _depth, item in payload)  # type: ignore[misc]
        elif tag == "table":
            header, rows = payload  # type: ignore[misc]
            parts.extend(header)
            for r in rows:
                parts.extend(r)
        elif tag == "code":
            parts.append(str(payload[1]))  # type: ignore[index]
        elif tag == "image":
            parts.append(str(payload[0]))  # type: ignore[index]
    return "\n".join(parts)


def build(md_path: str, out_path: str, kind: str = "detail", number_assets: bool = True) -> dict:
    """把一份 markdown 渲染成 docx，返回统计信息（供出稿工具汇总）。"""
    if kind not in KINDS:
        raise SystemExit(f"未知的版本类型 {kind!r}；可选：{' / '.join(KINDS)}")
    style = DocStyle.for_kind(kind)
    with open(md_path, "r", encoding="utf-8") as fh:
        text = fh.read()
    meta, blocks = parse_markdown(text)
    base_dir = os.path.dirname(os.path.abspath(md_path))

    document = Document()
    _set_default_style(document, style)
    footer_left = meta.source or (meta.paper.split("\n")[0].strip() if meta.paper else "")
    _page_setup(document, style, footer_left)

    if style.add_cover:
        _cover(document, meta)
    # 简略版与摘录版都**不生成标题行**（用户要求）：文件名即标题。

    if style.add_toc:
        _add_toc(document)

    counters = {"figure": 0, "table": 0}
    missing_images: list[str] = []
    for kind_tag, payload in blocks:
        if kind_tag == "heading":
            level, title = payload  # type: ignore[misc]
            level = max(1, min(4, level))
            p = document.add_paragraph(style=f"Heading {level}")
            size = style.h1_size if level == 1 else (style.h2_size if level == 2 else style.base_size)
            _run(p, title, size, bold=True, color=BLACK, cjk_font=FONT_HEI)
        elif kind_tag == "para":
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(style.base_size * 2)
            add_runs(p, str(payload), style.base_size)
        elif kind_tag == "ulist":
            for depth, item in payload:  # type: ignore[misc]
                p = document.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
                p.paragraph_format.space_after = Pt(style.body_space_after / 2)
                p.paragraph_format.line_spacing = style.line_spacing
                add_runs(p, item, style.base_size)
        elif kind_tag == "olist":
            for depth, item in payload:  # type: ignore[misc]
                p = document.add_paragraph(style="List Number" if depth == 0 else "List Number 2")
                p.paragraph_format.space_after = Pt(style.body_space_after / 2)
                p.paragraph_format.line_spacing = style.line_spacing
                add_runs(p, item, style.base_size)
        elif kind_tag == "quote":
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.3)
            _shade_paragraph(p, "F2F4F8")
            _run(p, "引文　", style.base_size - 0.5, italic=True, color=BLACK)
            add_runs(p, str(payload), style.base_size - 0.5)
        elif kind_tag == "code":
            _lang, body = payload  # type: ignore[misc]
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_after = Pt(6)
            _shade_paragraph(p, "F6F6F6")
            lines = str(body).split("\n")
            for idx, line in enumerate(lines):
                run = p.add_run(line)
                run.font.size = Pt(style.code_size)
                _bind(run)
                if idx != len(lines) - 1:
                    run.add_break()
        elif kind_tag == "table":
            header, rows = payload  # type: ignore[misc]
            if number_assets:
                counters["table"] += 1
                cap = document.add_paragraph()
                _run(cap, f"表 {counters['table']}", style.table_size, bold=True, color=BLACK)
            _add_table(document, header, rows, style)
        elif kind_tag == "image":
            caption, path = payload  # type: ignore[misc]
            probe = path if os.path.isabs(path) else os.path.join(base_dir, path)
            if not os.path.exists(probe) and not os.path.exists(os.path.abspath(path)):
                missing_images.append(path)
            if number_assets:
                counters["figure"] += 1
                caption = f"图 {counters['figure']}　{caption}" if caption else f"图 {counters['figure']}"
            _add_image(document, caption, path, style, base_dir)
        elif kind_tag == "hr":
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, "─" * 28, style.base_size - 1, color=BLACK)

    _refresh_fields_on_open(document)
    directory = os.path.dirname(os.path.abspath(out_path))
    if directory:
        os.makedirs(directory, exist_ok=True)
    _save_with_retry(document, out_path)

    plain = plain_text(blocks)
    return {
        "input": os.path.abspath(md_path),
        "output": os.path.abspath(out_path),
        "kind": kind,
        "blocks": len(blocks),
        "headings": sum(1 for t, _ in blocks if t == "heading"),
        "tables": sum(1 for t, _ in blocks if t == "table"),
        "images": sum(1 for t, _ in blocks if t == "image"),
        "quotes": sum(1 for t, _ in blocks if t == "quote"),
        "chars": len(plain),
        "cjk_chars": sum(1 for ch in plain if "\u4e00" <= ch <= "\u9fff"),
        "missing_images": missing_images,
        "toc": style.add_toc,
    }
